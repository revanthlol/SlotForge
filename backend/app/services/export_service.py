import csv
import io
import os
import uuid
from typing import Optional, List
from sqlalchemy.orm import Session

import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.models.timetable_version import TimetableVersion as VersionModel
from app.models.timetable_slot import TimetableSlot as SlotModel
from app.models.organization import Organization as OrgModel
from app.models.section import Section as SectionModel
from app.models.subject import Subject as SubjectModel
from app.models.teacher import Teacher as TeacherModel
from app.models.room import Room as RoomModel
from app.core.config import settings
from supabase import create_client

class ExportService:
    @staticmethod
    def get_sorted_slots(db: Session, version_id: uuid.UUID) -> List[any]:
        """
        Retrieves slots for the given timetable version, resolving names of teachers,
        rooms, subjects, and sections, sorted by day (Mon-Fri), period, and section name.
        """
        query = db.query(
            SlotModel.day,
            SlotModel.period,
            SectionModel.name.label("section_name"),
            SubjectModel.name.label("subject_name"),
            TeacherModel.name.label("teacher_name"),
            RoomModel.name.label("room_name")
        ).join(SectionModel, SlotModel.section_id == SectionModel.id)\
         .join(SubjectModel, SlotModel.subject_id == SubjectModel.id)\
         .join(TeacherModel, SlotModel.teacher_id == TeacherModel.id)\
         .join(RoomModel, SlotModel.room_id == RoomModel.id)\
         .filter(SlotModel.timetable_version_id == version_id)
         
        slots = query.all()
        
        # Define standard Day order for scheduling
        DAY_ORDER = {
            "Monday": 1,
            "Tuesday": 2,
            "Wednesday": 3,
            "Thursday": 4,
            "Friday": 5,
            "Saturday": 6,
            "Sunday": 7
        }
        
        return sorted(
            slots,
            key=lambda s: (DAY_ORDER.get(s.day, 99), s.period, s.section_name or "")
        )

    @staticmethod
    def generate_csv(slots: List[any]) -> bytes:
        """
        Generates CSV format of the timetable slots.
        """
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Day", "Period", "Section", "Subject", "Teacher", "Room"])
        for s in slots:
            writer.writerow([s.day, s.period, s.section_name, s.subject_name, s.teacher_name, s.room_name])
        return output.getvalue().encode("utf-8")

    @staticmethod
    def generate_excel(slots: List[any]) -> bytes:
        """
        Generates Excel format of the timetable slots.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Timetable"
        
        # Headers
        ws.append(["Day", "Period", "Section", "Subject", "Teacher", "Room"])
        
        # Data
        for s in slots:
            ws.append([s.day, s.period, s.section_name, s.subject_name, s.teacher_name, s.room_name])
            
        # Adjust Column Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
            
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @staticmethod
    def generate_pdf(slots: List[any], org_name: str, version_number: int, status: str) -> bytes:
        """
        Generates a beautifully styled ReportLab PDF of the timetable.
        """
        out = io.BytesIO()
        doc = SimpleDocTemplate(
            out,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )
        styles = getSampleStyleSheet()
        
        story = []
        
        # Header Styles
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#2d5d4f'), # Brand Pine Green
            spaceAfter=15
        )
        
        story.append(Paragraph(f"Timetable Report — {org_name}", title_style))
        story.append(Paragraph(f"Version: {version_number} | Status: {status.upper()}", styles['Normal']))
        story.append(Spacer(1, 15))
        
        # Table content
        table_data = [["Day", "Period", "Section", "Subject", "Teacher", "Room"]]
        for s in slots:
            table_data.append([
                s.day,
                str(s.period),
                s.section_name or "",
                s.subject_name or "",
                s.teacher_name or "",
                s.room_name or ""
            ])
            
        # Column widths: Day (80), Period (50), Section (80), Subject (120), Teacher (120), Room (90) -> Total 540 (fits exactly inside printable area)
        col_widths = [80, 50, 80, 120, 120, 90]
        
        t = Table(table_data, colWidths=col_widths)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2d5d4f')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('TOPPADDING', (0, 0), (-1, 0), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7f9f8')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e0e0e0')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
            ('TOPPADDING', (0, 1), (-1, -1), 6),
        ]))
        
        story.append(t)
        doc.build(story)
        return out.getvalue()

    @staticmethod
    def upload_file(org_id: uuid.UUID, filename: str, file_bytes: bytes) -> str:
        """
        Uploads the file to Supabase Storage exports bucket.
        Falls back to local static file serving if settings.DEV_MODE is True,
        if Supabase settings are placeholders, or if upload fails.
        """
        path = f"{org_id}/{filename}"
        
        is_placeholder = (
            "YOUR_PROJECT_REF" in settings.SUPABASE_URL or
            "your_publishable_key_here" in settings.SUPABASE_PUBLISHABLE_KEY or
            "your_secret_key_here" in settings.SUPABASE_SECRET_KEY
        )
        
        if not settings.DEV_MODE and not is_placeholder:
            try:
                supabase = create_client(settings.SUPABASE_URL, settings.SUPABASE_SECRET_KEY)
                
                # Upload file (with upsert Option)
                supabase.storage.from_("exports").upload(
                    path=path,
                    file=file_bytes,
                    file_options={"x-upsert": "true", "content-type": "application/octet-stream"}
                )
                
                # Retrieve signed URL valid for 1 hour
                res = supabase.storage.from_("exports").create_signed_url(path, 3600)
                if isinstance(res, dict) and "signedURL" in res:
                    return res["signedURL"]
                elif isinstance(res, dict) and "signed_url" in res:
                    return res["signed_url"]
                return str(res)
            except Exception as e:
                # Log or proceed to fallback
                pass
                
        # Local Static File Fallback
        static_dir = os.path.join("app", "static", "exports", str(org_id))
        os.makedirs(static_dir, exist_ok=True)
        
        local_filepath = os.path.join(static_dir, filename)
        with open(local_filepath, "wb") as f:
            f.write(file_bytes)
            
        # Return local static URL mapped in main.py
        return f"http://localhost:8000/static/exports/{org_id}/{filename}"
